"""Briques communes des vérifications Clean Cooking : anomalie, log cumulatif Excel
(Nouveau / Toujours ouvert / Résolu / Supprimé), rapport et email de confirmation.

Partagé par verify_inscription.py, verify_distribution.py, etc., pour qu'un correctif sur la
fusion du log profite à toutes les activités (même format de log : une ligne par anomalie,
identifiée par Dimension + Sous-dimension + Response Code + Description + ID ménage).
"""

import re
from collections import Counter
from datetime import datetime

from common.dates import parse_dt
from common.sharepoint import send_html_email

DIMENSIONS = ["Complétude", "Promptitude", "Validité", "Unicité", "Cohérence", "Fiabilité"]

LOG_HEADERS = [
    "Dimension", "Sous-dimension", "Response Code", "Household ID",
    "Description", "Détails", "Statut", "Première détection",
    "Dernière détection", "Date de résolution",
]

STATUTS_CLOS = ("Résolu", "Supprimé")


def val(r, col):
    return (r.get(col) or "").strip()


def is_digits(value):
    return re.fullmatch(r"[0-9]+", value) is not None


class Anomaly:
    def __init__(self, dimension, subdimension, response_code, household_id, description, details=""):
        self.dimension = dimension
        self.subdimension = subdimension
        self.response_code = response_code
        self.household_id = household_id
        self.description = description
        self.details = details


def anomaly_key(a):
    return (a.dimension, a.subdimension, a.response_code, a.description, a.household_id)


def has_entity(raw_responses, entity_type):
    """`raw_responses` : liste de réponses brutes partageant un même Response Code (voir
    common.mwater_client.fetch_raw_responses_by_code). Vrai si l'une référence une entité de ce
    type : le champ a été répondu même si le datagrid l'exporte vide (site supprimé/inaccessible)."""
    return any(e.get("entityType") == entity_type
               for resp in (raw_responses or []) for e in resp.get("entities", []))


def build_submitted_on_lookup(rows, code_col="Response Code", submitted_col="Submitted On"):
    lookup = {}
    for r in rows:
        dt = parse_dt(val(r, submitted_col))
        if val(r, code_col) and dt:
            lookup[val(r, code_col)] = dt.strftime("%d/%m/%Y")
    return lookup


def merge_with_log(current_anomalies, existing_rows, today_str, submitted_on_lookup=None,
                   existing_response_codes=None):
    """Même logique que carnet_de_bord/verify_carnet_de_bord.py : Nouveau, Toujours ouvert,
    Résolu (date = Submitted On de la réponse), Supprimé (réponse mWater disparue)."""
    submitted_on_lookup = submitted_on_lookup or {}
    existing_response_codes = existing_response_codes or set()
    existing_open, already_closed = {}, []
    for row in existing_rows:
        if row.get("Statut") in STATUTS_CLOS:
            already_closed.append(row)
        else:
            key = (row.get("Dimension"), row.get("Sous-dimension"), row.get("Response Code"),
                   row.get("Description"), row.get("Household ID"))
            existing_open[key] = row

    merged, seen_keys, new_count = [], set(), 0
    for a in current_anomalies:
        key = anomaly_key(a)
        seen_keys.add(key)
        existing = existing_open.get(key)
        if not existing:
            new_count += 1
        merged.append({
            "Dimension": a.dimension, "Sous-dimension": a.subdimension, "Response Code": a.response_code,
            "Household ID": a.household_id, "Description": a.description, "Détails": a.details,
            "Statut": "Toujours ouvert" if existing else "Nouveau",
            "Première détection": existing.get("Première détection") if existing else today_str,
            "Dernière détection": today_str, "Date de résolution": "",
        })

    resolved_count = deleted_count = 0
    for key, row in existing_open.items():
        if key in seen_keys:
            continue
        row = dict(row)
        if row.get("Response Code") not in existing_response_codes:
            row["Statut"], row["Date de résolution"] = "Supprimé", row.get("Première détection", "")
            deleted_count += 1
        else:
            row["Statut"] = "Résolu"
            row["Date de résolution"] = submitted_on_lookup.get(row.get("Response Code"), today_str)
            resolved_count += 1
        merged.append(row)

    merged.extend(already_closed)
    return merged, new_count, resolved_count, deleted_count


def build_log_report(merged_rows, output_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Anomalies"
    ws.append(LOG_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in merged_rows:
        ws.append([row.get(h, "") for h in LOG_HEADERS])

    open_rows = [r for r in merged_rows if r.get("Statut") not in STATUTS_CLOS]
    summary = wb.create_sheet("Résumé")
    summary.append(["Dimension", "Nombre d'anomalies ouvertes"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    counts = Counter(r.get("Dimension") for r in open_rows)
    for dimension in DIMENSIONS:
        summary.append([dimension, counts.get(dimension, 0)])
    summary.append(["Total ouvert", len(open_rows)])
    wb.save(output_path)


def send_confirmation_email(token, sender, recipients, activity_label, file_name, open_count,
                            new_count, resolved_count, deleted_count, counts_by_dimension):
    lines = "".join(f"<li>{dim} : {n}</li>" for dim, n in counts_by_dimension.items())
    body_html = (
        f"<p>Vérification de données \"{activity_label}\" exécutée.</p>"
        f"<p>Anomalies actuellement ouvertes : <b>{open_count}</b> (dont {new_count} nouvelles) — "
        f"{resolved_count} résolue(s) et {deleted_count} supprimée(s) (réponse mWater disparue) "
        "depuis la dernière exécution.</p>"
        f"<ul>{lines}</ul>"
        f"<p>Log mis à jour sur SharePoint : <b>{file_name}</b></p>"
    )
    send_html_email(token, sender, recipients, f"Vérification de données — {file_name}", body_html)


def print_summary(anomalies):
    detail = Counter((a.dimension, a.subdimension) for a in anomalies)
    for (dimension, sub), n in sorted(detail.items(), key=lambda kv: (DIMENSIONS.index(kv[0][0]), -kv[1])):
        print(f"    {dimension} / {sub} : {n}")


def today():
    return datetime.now().strftime("%d/%m/%Y")
