#!/usr/bin/env python3
"""
Vérification de données — Clean Cooking : Inscription
======================================================

Applique les six dimensions du "Manuel de vérification — Clean Cooking : Inscription"
(Complétude, Promptitude, Validité, Unicité, Cohérence, Fiabilité) au datagrid mWater
"Inscription des utilisateurs des foyers améliorés MAGNASOA" (formulaire
434df1bc4f144644b7ba52257ba30e2a), puis met à jour un log Excel cumulatif
(data_verification_clean_cooking_inscription_log.xlsx : Nouveau / Toujours ouvert / Résolu /
Supprimé) déposé sur SharePoint, et envoie un email de confirmation.

Mode test : avec DRY_RUN=true, le script détecte et affiche le décompte des anomalies par
dimension et sous-dimension, sans toucher à SharePoint ni envoyer d'email (aucun secret
Azure/email n'est alors nécessaire).

Écarts assumés avec le guide (à valider) :
- Cohérence village / adresse : le guide demande que le code village corresponde à l'adresse
  (région, district, commune). Implémenté ici comme "un village ne doit pas être rattaché à
  plusieurs triplets région/district/commune" ; les lignes minoritaires sont signalées.
  La phrase "une même adresse ne doit pas pointer vers deux villages" n'est pas appliquée telle
  quelle : une commune contient légitimement plusieurs villages.
- Cohérence ménage de référence : le datagrid exporte l'ID ménage vide si le site lié n'existe
  plus ; ce cas est couvert par la Complétude (distinction "site supprimé/inaccessible").
"""

import os
import re
from collections import Counter, defaultdict
from datetime import datetime

from common.dates import parse_dt
from common.mwater_client import download_datagrid_with_retry, fetch_raw_responses_by_code, mwater_login
from common.sharepoint import (
    download_existing_log,
    folder_link_from_env,
    graph_token,
    resolve_share_link,
    send_html_email,
    upload_to_sharepoint,
)

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

DATAGRID_INSCRIPTION = "ac3ae2ce1f7540f586ae1ab5673419a8"
FORM_INSCRIPTION = "434df1bc4f144644b7ba52257ba30e2a"

LOG_FILE_NAME = "data_verification_clean_cooking_inscription_log.xlsx"

LOG_HEADERS = [
    "Dimension", "Sous-dimension", "Response Code", "Household ID",
    "Description", "Détails", "Statut", "Première détection",
    "Dernière détection", "Date de résolution",
]

DIMENSIONS = ["Complétude", "Promptitude", "Validité", "Unicité", "Cohérence", "Fiabilité"]

# Noms de colonnes du datagrid (orthographe exacte du portail, fautes comprises)
C_STATUS = "Status"
C_CODE = "Response Code"
C_HOUSEHOLD = "Household ID code > Unique ID"
C_VILLAGE = "Code village > Unique ID"
C_REPRESENTANT = "Name of household representative"
C_CIN = "Numéro de carte d'identité nationale (CIN) du représentant du ménage"
C_PHONE = "Telephone number (if available) of the household"
C_PHONE_NA = "Telephone number (if available) of the household (Not Applicable)"
C_ADRESSE = {k: f"Adresse du ménage > {k}" for k in ("REGION", "DISTRICT", "COMMUNE", "FOKONTANY")}
C_MODELE = "Modèle/type de foyer amélioré commandé"
C_QUANTITE = "Quantité de foyer améliorés efficaces non fonctionnels à remettre"
C_MODE = "Mode d'utilisation"
C_PAIEMENT = "Etat de paiment"
C_RECU = "Annexe G : Reçu de paiement"
C_DRAFTED = "Drafted On"
C_SUBMITTED = "Submitted On"
C_REJECTION = "Rejection message"

PIECES_JUSTIFICATIVES = {
    "Photo de la maison": "Photo de la maison",
    "Photo de la carte d'identité": "Photo carte d'identité nationale",
    "Fiche d'inscription (Annexe A)": "Annexe A : Fiche d'inscription des membres",
    "Accord carbone (Annexe B)": "Annexe B : Accord Carbone",
    "Consentement (Annexe D)": "Annexe D : Formulaire de consentement pour la publication de données et de photos",
}

MODELES_VALIDES = {
    "ICS Bois de chauffe : Type 2 Ecoa Wood G3.1 (Burn)",
    "ICS Charbon : Type 1 Jikokoa Plus Stainles (Burn)",
    'ICS Charbon : Type 3 Jikokoa Plus Stainles (Burn) (Ecoa en charbon)" le ménage utilise un seul foyer au Charbon"',
}
MODES_VALIDES = {"Domestique", "Commerciale", "Institutionnelle"}
PAIEMENTS_VALIDES = {"Non payé", "Payé totalement"}
PAIEMENT_TOTAL = "Payé totalement"

STATUTS_CLOS = ("Résolu", "Supprimé")


def val(r, col):
    return (r.get(col) or "").strip()


def is_digits(value):
    return re.fullmatch(r"[0-9]+", value) is not None


# ---------------------------------------------------------------------------
# Anomalie : structure commune
# ---------------------------------------------------------------------------

class Anomaly:
    def __init__(self, dimension, subdimension, response_code, household_id, description, details=""):
        self.dimension = dimension
        self.subdimension = subdimension
        self.response_code = response_code
        self.household_id = household_id
        self.description = description
        self.details = details


def new(dimension, subdimension, r, description, details=""):
    return Anomaly(dimension, subdimension, val(r, C_CODE), val(r, C_HOUSEHOLD), description, details)


# ---------------------------------------------------------------------------
# Entités liées supprimées/inaccessibles (distinction Complétude)
# ---------------------------------------------------------------------------

def has_entity(raw_responses, entity_type):
    """`raw_responses` : liste de réponses brutes partageant un même Response Code (voir
    fetch_raw_responses_by_code)."""
    return any(e.get("entityType") == entity_type
               for resp in (raw_responses or []) for e in resp.get("entities", []))


def resolve_missing_entities(rows, client_id):
    """Pour les lignes dont l'ID ménage ou le code village est vide dans le datagrid, regarde la
    réponse brute mWater : si elle référence bien une entité (ménage / village), le champ a été
    répondu mais le site lié est supprimé ou inaccessible -> pas un champ jamais renseigné."""
    codes = sorted({val(r, C_CODE) for r in rows if not val(r, C_HOUSEHOLD) or not val(r, C_VILLAGE)})
    if not codes:
        return {}
    raw = fetch_raw_responses_by_code(client_id, FORM_INSCRIPTION, codes)
    return {c: {"household": has_entity(raw.get(c), "household"),
                "community": has_entity(raw.get(c), "community")} for c in codes}


# ---------------------------------------------------------------------------
# Dimensions
# ---------------------------------------------------------------------------

def check_completude(rows, entities):
    out = []
    for r in rows:
        status = val(r, C_STATUS)
        if status != "Final":
            out.append(new("Complétude", "Réponse non finalisée", r,
                           "Inscription non finalisée", f"Status : {status or 'vide'}"))

        flags = entities.get(val(r, C_CODE), {})
        if not val(r, C_HOUSEHOLD):
            if flags.get("household"):
                out.append(new("Complétude", "Ménage supprimé/inaccessible", r,
                               "Household ID code répondu mais site associé supprimé ou inaccessible",
                               "Non résolu dans l'export mWater malgré une réponse valide (à vérifier dans le portail)"))
            else:
                out.append(new("Complétude", "Identifiant du ménage manquant", r, "Household ID code vide"))
        if not val(r, C_VILLAGE):
            if flags.get("community"):
                out.append(new("Complétude", "Village supprimé/inaccessible", r,
                               "Code village répondu mais site associé supprimé ou inaccessible",
                               "Non résolu dans l'export mWater malgré une réponse valide (à vérifier dans le portail)"))
            else:
                out.append(new("Complétude", "Village manquant", r, "Code village vide"))

        champs = {
            "Représentant du ménage": C_REPRESENTANT,
            "Modèle de foyer commandé": C_MODELE,
            "Mode d'utilisation": C_MODE,
            "État de paiement": C_PAIEMENT,
        }
        for label, col in champs.items():
            if not val(r, col):
                out.append(new("Complétude", "Champ obligatoire vide", r, f"{label} non renseigné"))
        manquants = [k for k, col in C_ADRESSE.items() if not val(r, col)]
        if manquants:
            out.append(new("Complétude", "Champ obligatoire vide", r, "Adresse du ménage incomplète",
                           "Niveaux vides : " + ", ".join(manquants)))
        if not val(r, C_PHONE) and val(r, C_PHONE_NA).lower() != "true":
            out.append(new("Complétude", "Champ obligatoire vide", r,
                           "Téléphone non renseigné (« non applicable » non coché)"))

        for label, col in PIECES_JUSTIFICATIVES.items():
            if not val(r, col):
                out.append(new("Complétude", "Pièce justificative manquante", r, f"Pièce manquante : {label}"))
        if val(r, C_PAIEMENT) == PAIEMENT_TOTAL and not val(r, C_RECU):
            out.append(new("Complétude", "Reçu de paiement manquant", r,
                           "Reçu de paiement (Annexe G) manquant alors que l'état de paiement est « Payé totalement »"))
    return out


def check_promptitude(rows):
    out = []
    for r in rows:
        drafted, submitted = parse_dt(val(r, C_DRAFTED)), parse_dt(val(r, C_SUBMITTED))
        if drafted and submitted and submitted < drafted:
            out.append(new("Promptitude", "Chronologie Drafted/Submitted", r,
                           "Submitted On antérieur à Drafted On", f"Drafted On : {drafted} / Submitted On : {submitted}"))
    return out


def check_validite(rows):
    out = []
    for r in rows:
        for label, col in (("Household ID code", C_HOUSEHOLD), ("Code village", C_VILLAGE),
                           ("Numéro de CIN", C_CIN), ("Téléphone", C_PHONE)):
            v = val(r, col)
            if v and not is_digits(v):
                out.append(new("Validité", f"Format {label}", r, f"{label} non numérique", f"Valeur : '{v}'"))
        for label, col, valides in (("Modèle de foyer commandé", C_MODELE, MODELES_VALIDES),
                                    ("Mode d'utilisation", C_MODE, MODES_VALIDES),
                                    ("État de paiement", C_PAIEMENT, PAIEMENTS_VALIDES)):
            v = val(r, col)
            if v and v not in valides:
                out.append(new("Validité", f"Valeur {label}", r, f"{label} hors des valeurs du formulaire",
                               f"Valeur : '{v}'"))
        q = val(r, C_QUANTITE)
        if q and not is_digits(q):
            out.append(new("Validité", "Quantité de foyers à remettre", r,
                           "Quantité de foyers non fonctionnels à remettre invalide (entier ≥ 0 attendu)",
                           f"Valeur : '{q}'"))
    return out


def check_unicite(rows):
    out = []
    counts = Counter(val(r, C_HOUSEHOLD) for r in rows if val(r, C_STATUS) != "Draft" and val(r, C_HOUSEHOLD))
    for r in rows:
        household = val(r, C_HOUSEHOLD)
        if val(r, C_STATUS) != "Draft" and household and counts[household] > 1:
            out.append(new("Unicité", "Doublon ménage", r, "Ménage inscrit plusieurs fois",
                           f"Household ID code {household} présent dans {counts[household]} réponses"))
        elif val(r, C_STATUS) == "Rejected" and "doublon" in val(r, C_REJECTION).lower():
            out.append(new("Unicité", "Rejet pour doublon", r, "Rejet mWater pour doublon", val(r, C_REJECTION)))
    return out


def check_coherence(rows):
    """Un code village doit correspondre à un seul triplet région/district/commune."""
    triplets = defaultdict(Counter)
    for r in rows:
        village = val(r, C_VILLAGE)
        triplet = tuple(val(r, C_ADRESSE[k]) for k in ("REGION", "DISTRICT", "COMMUNE"))
        if village and all(triplet):
            triplets[village][triplet] += 1
    out = []
    for r in rows:
        village = val(r, C_VILLAGE)
        triplet = tuple(val(r, C_ADRESSE[k]) for k in ("REGION", "DISTRICT", "COMMUNE"))
        if village in triplets and all(triplet) and len(triplets[village]) > 1:
            majoritaire = sorted(triplets[village].items(), key=lambda kv: (-kv[1], kv[0]))[0][0]
            if triplet != majoritaire:
                out.append(new("Cohérence", "Village / adresse", r,
                               "Adresse (région/district/commune) incohérente avec le village",
                               f"Village {village} : adresse saisie {' / '.join(triplet)} ; "
                               f"adresse la plus fréquente pour ce village {' / '.join(majoritaire)}"))
    return out


def check_fiabilite(rows):
    out = []
    for r in rows:
        # Un rejet n'est une anomalie que si la réponse est toujours au statut Rejected.
        if val(r, C_STATUS) == "Rejected":
            out.append(new("Fiabilité", "Rejet mWater non corrigé", r, "Réponse rejetée non corrigée",
                           val(r, C_REJECTION) or "Aucun message de rejet"))
    return out


# ---------------------------------------------------------------------------
# Log cumulatif : Nouveau / Toujours ouvert / Résolu / Supprimé
# ---------------------------------------------------------------------------

def anomaly_key(a):
    return (a.dimension, a.subdimension, a.response_code, a.description, a.household_id)


def build_submitted_on_lookup(rows):
    lookup = {}
    for r in rows:
        dt = parse_dt(val(r, C_SUBMITTED))
        if val(r, C_CODE) and dt:
            lookup[val(r, C_CODE)] = dt.strftime("%d/%m/%Y")
    return lookup


def merge_with_log(current_anomalies, existing_rows, today_str, submitted_on_lookup=None,
                   existing_response_codes=None):
    """Même logique que carnet_de_bord/verify_carnet_de_bord.py : Nouveau, Toujours ouvert,
    Résolu (date = Submitted On de la réponse), Supprimé (réponse mWater disparue)."""
    submitted_on_lookup = submitted_on_lookup or {}
    existing_response_codes = existing_response_codes or set()
    existing_open, already_closed = {}, []
    for row in existing_rows:
        if row.get("Statut") in STATUTS_CLOS:
            already_closed.append(row)
        else:
            key = (row.get("Dimension"), row.get("Sous-dimension"), row.get("Response Code"),
                   row.get("Description"), row.get("Household ID"))
            existing_open[key] = row

    merged, seen_keys, new_count = [], set(), 0
    for a in current_anomalies:
        key = anomaly_key(a)
        seen_keys.add(key)
        existing = existing_open.get(key)
        if not existing:
            new_count += 1
        merged.append({
            "Dimension": a.dimension, "Sous-dimension": a.subdimension, "Response Code": a.response_code,
            "Household ID": a.household_id, "Description": a.description, "Détails": a.details,
            "Statut": "Toujours ouvert" if existing else "Nouveau",
            "Première détection": existing.get("Première détection") if existing else today_str,
            "Dernière détection": today_str, "Date de résolution": "",
        })

    resolved_count = deleted_count = 0
    for key, row in existing_open.items():
        if key in seen_keys:
            continue
        row = dict(row)
        if row.get("Response Code") not in existing_response_codes:
            row["Statut"], row["Date de résolution"] = "Supprimé", row.get("Première détection", "")
            deleted_count += 1
        else:
            row["Statut"] = "Résolu"
            row["Date de résolution"] = submitted_on_lookup.get(row.get("Response Code"), today_str)
            resolved_count += 1
        merged.append(row)

    merged.extend(already_closed)
    return merged, new_count, resolved_count, deleted_count


def build_log_report(merged_rows, output_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Anomalies"
    ws.append(LOG_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in merged_rows:
        ws.append([row.get(h, "") for h in LOG_HEADERS])

    open_rows = [r for r in merged_rows if r.get("Statut") not in STATUTS_CLOS]
    summary = wb.create_sheet("Résumé")
    summary.append(["Dimension", "Nombre d'anomalies ouvertes"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    counts = Counter(r.get("Dimension") for r in open_rows)
    for dimension in DIMENSIONS:
        summary.append([dimension, counts.get(dimension, 0)])
    summary.append(["Total ouvert", len(open_rows)])
    wb.save(output_path)


def send_confirmation_email(token, sender, recipients, file_name, open_count, new_count,
                            resolved_count, deleted_count, counts_by_dimension):
    lines = "".join(f"<li>{dim} : {n}</li>" for dim, n in counts_by_dimension.items())
    body_html = (
        "<p>Vérification de données \"Clean Cooking — Inscription\" exécutée.</p>"
        f"<p>Anomalies actuellement ouvertes : <b>{open_count}</b> (dont {new_count} nouvelles) — "
        f"{resolved_count} résolue(s) et {deleted_count} supprimée(s) (réponse mWater disparue) "
        "depuis la dernière exécution.</p>"
        f"<ul>{lines}</ul>"
        f"<p>Log mis à jour sur SharePoint : <b>{file_name}</b></p>"
    )
    send_html_email(token, sender, recipients, f"Vérification de données — {file_name}", body_html)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    dry_run = os.environ.get("DRY_RUN", "").strip().lower() in ("1", "true", "yes")

    print("Authentification mWater...")
    client_id = mwater_login(os.environ["MWATER_USERNAME"].strip(), os.environ["MWATER_PASSWORD"].strip())

    print("Téléchargement du datagrid Inscription...")
    # ~90 000 lignes : mWater répond parfois 524 (export trop long à générer) -> plus de tentatives
    rows = download_datagrid_with_retry(DATAGRID_INSCRIPTION, client_id, attempts=8, wait_seconds=60)
    print(f"  {len(rows)} lignes")

    print("Résolution des ID ménage / villages vides (site supprimé ou jamais répondu)...")
    entities = resolve_missing_entities(rows, client_id)
    print(f"  {len(entities)} réponses vérifiées sur l'API brute")

    print("Application des règles de vérification (6 dimensions)...")
    anomalies = []
    anomalies += check_completude(rows, entities)
    anomalies += check_promptitude(rows)
    anomalies += check_validite(rows)
    anomalies += check_unicite(rows)
    anomalies += check_coherence(rows)
    anomalies += check_fiabilite(rows)
    print(f"  {len(anomalies)} anomalies détectées")
    detail = Counter((a.dimension, a.subdimension) for a in anomalies)
    for (dimension, sub), n in sorted(detail.items(), key=lambda kv: (DIMENSIONS.index(kv[0][0]), -kv[1])):
        print(f"    {dimension} / {sub} : {n}")

    if dry_run:
        print("DRY_RUN : aucun accès SharePoint, aucun email, aucun log écrit.")
        return

    azure_tenant_id = os.environ["AZURE_TENANT_ID"].strip()
    azure_client_id = os.environ["AZURE_CLIENT_ID"].strip()
    azure_client_secret = os.environ["AZURE_CLIENT_SECRET"].strip()
    sharepoint_folder_link = folder_link_from_env("SHAREPOINT_FOLDER_LINK_INSCRIPTION")
    email_sender = os.environ["EMAIL_SENDER"].strip()
    email_recipients = os.environ["EMAIL_RECIPIENTS"].strip()

    print("Authentification Microsoft Graph...")
    token = graph_token(azure_tenant_id, azure_client_id, azure_client_secret)
    print("Résolution du lien SharePoint...")
    drive_id, folder_item_id = resolve_share_link(token, sharepoint_folder_link)

    print("Téléchargement du log existant sur SharePoint...")
    existing_rows = download_existing_log(token, drive_id, folder_item_id, LOG_FILE_NAME)
    print(f"  {len(existing_rows)} lignes déjà présentes dans le log")

    today_str = datetime.now().strftime("%d/%m/%Y")
    merged_rows, new_count, resolved_count, deleted_count = merge_with_log(
        anomalies, existing_rows, today_str,
        submitted_on_lookup=build_submitted_on_lookup(rows),
        existing_response_codes={val(r, C_CODE) for r in rows})
    open_rows = [r for r in merged_rows if r.get("Statut") not in STATUTS_CLOS]
    print(f"  {len(open_rows)} anomalies ouvertes ({new_count} nouvelles, {resolved_count} résolues, "
          f"{deleted_count} supprimées)")

    output_path = f"/tmp/{LOG_FILE_NAME}"
    build_log_report(merged_rows, output_path)
    print("Dépôt du log mis à jour sur SharePoint...")
    upload_to_sharepoint(token, drive_id, folder_item_id, output_path, LOG_FILE_NAME)

    print("Envoi de l'email de confirmation...")
    send_confirmation_email(token, email_sender, email_recipients, LOG_FILE_NAME, len(open_rows),
                            new_count, resolved_count, deleted_count,
                            Counter(r.get("Dimension") for r in open_rows))
    print("Terminé.")


if __name__ == "__main__":
    main()
