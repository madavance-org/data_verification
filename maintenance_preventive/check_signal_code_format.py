#!/usr/bin/env python3
"""
Vérification ponctuelle — Format des Signal code (Maintenance préventive)
==========================================================================

Contrôle indépendant de la Validité "Signal reference" déjà faite dans
appel_maintenance_preventive/verify_maintenance_preventive.py (qui ne porte
que sur Appel maintenance préventive) : ici, on vérifie le champ "Signal
code" tel que saisi directement dans le formulaire Maintenance préventive,
pour repérer les codes qui ne respectent pas le format
{DEPLOYMENT}_{JJMMAAAA}_{E|S}{N} (ex. MAR_21072026_E14).

Même principe de log que verify_maintenance_preventive.py, mais dans un
fichier séparé (data_verification_signal_code_maintenance_log.xlsx) : seules
les anomalies sont journalisées, avec un suivi Nouveau / Toujours ouvert /
Résolu à travers les exécutions successives (action déclenchée
manuellement, pas de cron).

Pour les erreurs de formatage simples (espace, tiret/slash au lieu
d'underscore, underscore manquant avant E/S, E/S et numéro inversés,
caractère parasite en fin de code), une valeur corrigée est proposée
automatiquement dans la colonne "Signal code correct". Si l'anomalie
est ambiguë (chiffres manquants/en trop dans la date, code trop
dégradé), la colonne reste vide : correction manuelle nécessaire.
"""

import os
import re
from collections import Counter
from datetime import datetime

from common.dates import parse_dt
from common.mwater_client import download_datagrid, mwater_login
from common.sharepoint import download_existing_log, graph_token, resolve_share_link, upload_to_sharepoint

DATAGRID_MAINTENANCE = "2cfe0ba7ac264d119cfc8964b5f3cebc"

LOG_FILE_NAME = "data_verification_signal_code_maintenance_log.xlsx"

LOG_HEADERS = [
    "Formulaire", "Response Code", "Signal code", "Anomalie",
    "Signal code correct", "Statut", "Première détection",
    "Dernière détection", "Date de résolution",
]

SIGNAL_CODE_RE = re.compile(r"^([A-Z]+)_(\d{2})(\d{2})(\d{4})_([ES])(\d+)$")

# Format "propre" mais avec séparateurs/ordre incorrects : ex. "MAR-08052026-E1",
# "MAR _09072026_E2", "MAR_26/02/2026_S14", "FTU_28042026E1"
LOOSE_RE = re.compile(
    r"^\s*([A-Za-z]+)[\s_\-]*(\d{2})[\s/\-]?(\d{2})[\s/\-]?(\d{4})"
    r"[\s_\-]*([EeSs])[\s_\-]*(\d+)\.?\s*$"
)

# Même chose mais avec E/S et numéro inversés : ex. "MAR_11022025_1E"
SWAPPED_RE = re.compile(
    r"^\s*([A-Za-z]+)[\s_\-]*(\d{2})[\s/\-]?(\d{2})[\s/\-]?(\d{4})"
    r"[\s_\-]*(\d+)[\s_\-]*([EeSs])\.?\s*$"
)


# ---------------------------------------------------------------------------
# Correction proposée
# ---------------------------------------------------------------------------

def propose_correction(raw):
    """Tente de reconstruire un Signal code valide à partir d'une valeur mal formatée.
    Retourne (valeur_corrigee_ou_None, note). None si le code est trop ambigu/dégradé
    pour une correction automatique fiable (nécessite une vérification manuelle)."""

    m = LOOSE_RE.match(raw)
    if not m:
        m = SWAPPED_RE.match(raw)
        if m:
            prefix, dd, mm, yyyy, num, es = m.groups()
        else:
            return None, "Format non reconnu automatiquement"
    else:
        prefix, dd, mm, yyyy, es, num = m.groups()

    # Sanité de la date avant de proposer une correction
    try:
        datetime(int(yyyy), int(mm), int(dd))
    except ValueError:
        return None, "Date invalide dans le code (jour/mois incohérents)"

    corrected = f"{prefix.upper()}_{dd}{mm}{yyyy}_{es.upper()}{num}"
    if corrected == raw:
        return None, "Format non reconnu automatiquement"
    return corrected, ""


def check_format(rows, form_label):
    """Retourne la liste des anomalies (dict) pour un formulaire donné."""
    anomalies = []
    for r in rows:
        code = (r.get("Signal code") or "").strip()
        if not code or SIGNAL_CODE_RE.match(code):
            continue  # vide ou déjà conforme : pas une anomalie

        corrected, note = propose_correction(code)
        description = "Signal code ne respecte pas le format attendu"
        if note:
            description += f" ({note})"

        anomalies.append({
            "Formulaire": form_label,
            "Response Code": r.get("Response Code", ""),
            "Signal code": code,
            "Anomalie": description,
            "Signal code correct": corrected or "",
            "Status brut mWater": r.get("Status", ""),
            "Submitted On": r.get("Submitted On", ""),
        })
    return anomalies


# ---------------------------------------------------------------------------
# Log persistant — Nouveau / Toujours ouvert / Résolu (même principe que
# verify_maintenance_preventive.py)
# ---------------------------------------------------------------------------

def anomaly_key(a):
    return (a["Formulaire"], a["Response Code"])


def merge_with_log(current_anomalies, existing_rows, today_str):
    existing_open = {}
    already_resolved = []
    for row in existing_rows:
        key = (row.get("Formulaire"), row.get("Response Code"))
        if row.get("Statut") == "Résolu":
            already_resolved.append(row)
        else:
            existing_open[key] = row

    merged = []
    seen_keys = set()
    new_count = 0

    for a in current_anomalies:
        key = anomaly_key(a)
        seen_keys.add(key)
        existing = existing_open.get(key)
        if not existing:
            new_count += 1
        merged.append({
            "Formulaire": a["Formulaire"],
            "Response Code": a["Response Code"],
            "Signal code": a["Signal code"],
            "Anomalie": a["Anomalie"],
            "Signal code correct": a["Signal code correct"],
            "Statut": "Toujours ouvert" if existing else "Nouveau",
            "Première détection": existing.get("Première détection") if existing else today_str,
            "Dernière détection": today_str,
            "Date de résolution": "",
        })

    resolved_count = 0
    submitted_on_lookup = {
        a["Response Code"]: parse_dt(a["Submitted On"]).strftime("%d/%m/%Y")
        for a in current_anomalies if parse_dt(a["Submitted On"])
    }
    for key, row in existing_open.items():
        if key not in seen_keys:
            row = dict(row)
            row["Statut"] = "Résolu"
            row["Date de résolution"] = submitted_on_lookup.get(row.get("Response Code"), today_str)
            merged.append(row)
            resolved_count += 1

    merged.extend(already_resolved)
    return merged, new_count, resolved_count


def build_log_report(merged_rows, output_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Anomalies"
    ws.append(LOG_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in merged_rows:
        ws.append([row.get(h, "") for h in LOG_HEADERS])

    open_rows = [r for r in merged_rows if r.get("Statut") != "Résolu"]
    summary = wb.create_sheet("Résumé")
    summary.append(["Formulaire", "Nombre d'anomalies ouvertes"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    counts = Counter(r.get("Formulaire") for r in open_rows)
    summary.append(["Maintenance préventive", counts.get("Maintenance préventive", 0)])
    summary.append(["Total ouvert", len(open_rows)])

    wb.save(output_path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    mwater_username = os.environ["MWATER_USERNAME"].strip()
    mwater_password = os.environ["MWATER_PASSWORD"].strip()

    azure_tenant_id = os.environ["AZURE_TENANT_ID"].strip()
    azure_client_id = os.environ["AZURE_CLIENT_ID"].strip()
    azure_client_secret = os.environ["AZURE_CLIENT_SECRET"].strip()

    sharepoint_folder_link = os.environ["SHAREPOINT_FOLDER_LINK"].strip()

    print("Authentification mWater...")
    client_id = mwater_login(mwater_username, mwater_password)

    print("Téléchargement du datagrid Maintenance préventive...")
    maintenance_rows = download_datagrid(DATAGRID_MAINTENANCE, client_id)
    print(f"  Maintenance préventive : {len(maintenance_rows)} réponses")

    print("Vérification du format des Signal code...")
    anomalies = check_format(maintenance_rows, "Maintenance préventive")
    print(f"  {len(anomalies)} anomalies détectées")

    print("Authentification Microsoft Graph...")
    token = graph_token(azure_tenant_id, azure_client_id, azure_client_secret)

    print("Résolution du lien SharePoint...")
    sharepoint_drive_id, sharepoint_folder_item_id = resolve_share_link(token, sharepoint_folder_link)

    print("Téléchargement du log existant sur SharePoint...")
    existing_rows = download_existing_log(token, sharepoint_drive_id, sharepoint_folder_item_id, LOG_FILE_NAME,
                                           key_column="Formulaire")
    print(f"  {len(existing_rows)} lignes déjà présentes dans le log")

    today_str = datetime.now().strftime("%d/%m/%Y")
    merged_rows, new_count, resolved_count = merge_with_log(anomalies, existing_rows, today_str)
    open_rows = [r for r in merged_rows if r.get("Statut") != "Résolu"]
    print(f"  {len(open_rows)} anomalies ouvertes ({new_count} nouvelles, {resolved_count} résolues)")

    output_path = f"/tmp/{LOG_FILE_NAME}"
    build_log_report(merged_rows, output_path)
    print(f"Rapport généré : {output_path}")

    print("Dépôt du log mis à jour sur SharePoint...")
    upload_to_sharepoint(token, sharepoint_drive_id, sharepoint_folder_item_id, output_path, LOG_FILE_NAME)

    print("Terminé.")


if __name__ == "__main__":
    main()
