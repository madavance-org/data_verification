#!/usr/bin/env python3
"""
Vérification de données — Maintenance préventive régulière
==========================================================

Applique les six dimensions du "Manuel de vérification de données" MadAvance à
l'activité "Maintenance préventive régulière" (formulaire
de26d89a5c8a4452b42158c622be20d0, datagrid DATAGRID_MAINTENANCE). Règles : voir
Manuel_Maintenance_preventive_reguliere (ClickUp), qui fait foi.

Particularités par rapport aux autres scripts de ce repo :
  - Les règles croisées avec Appel maintenance préventive (signal code présent dans le
    bon formulaire, date du signal code <= date de complétion) ne sont PAS refaites ici :
    elles sont déjà vérifiées côté Appel.
  - Les noms de colonnes du datagrid sont résolus à l'exécution (voir find_column) : si
    une colonne attendue est introuvable, les contrôles qui en dépendent sont sautés avec
    un avertissement qui liste les colonnes réellement disponibles, plutôt que de produire
    des anomalies factices.
  - Fiabilité : la liste des points d'eau réhabilités avec succès est construite depuis les
    réponses brutes de l'ancien formulaire combiné (questions de la section « First
    rehabilitation », pas celles de la section « Réparation après panne », qui portent
    des libellés presque identiques) — voir build_success_list.
  - Log dédié (data_verification_maintenance_reguliere_log.xlsx), distinct des autres
    activités ; écriture parallèle dans le registre mWater via common/mwater_registry.py.
"""

import json
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime

import requests

from common.dates import parse_dt
from common.http_utils import raise_for_status_verbose
from common.mwater_client import (
    MWATER_API_BASE,
    download_datagrid,
    extract_water_point_code,
    fetch_raw_responses_by_code,
    mwater_login,
)
from common.mwater_registry import insert_log_in_mwater
from common.sharepoint import (
    download_existing_log,
    graph_token,
    resolve_share_link,
    send_html_email,
    upload_to_sharepoint,
)
from common.signal_code import SIGNAL_CODE_RE, propose_correction

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

DATAGRID_MAINTENANCE = "2cfe0ba7ac264d119cfc8964b5f3cebc"
FORM_MAINTENANCE = "de26d89a5c8a4452b42158c622be20d0"

# Ancien formulaire combiné : sert uniquement à construire la liste des points d'eau
# réhabilités avec succès (dimension Fiabilité). IDs de questions de la section
# « First rehabilitation > Repairable » (vérifiés dans la définition du formulaire le
# 13/10/2026) ; les questions « Réparation après panne » de ce même formulaire ont des
# IDs différents et ne doivent pas être utilisées ici.
FORM_COMBINE = "86cf66efdd3749dd8a121314bab3675a"
Q_TYPE_TRAVAUX = "7f78d719b9f242d8886df7bb88640a80"
CHOICE_PREMIERE_REHAB = "DQcV1NT"
Q_REPARATION_REUSSIE = "931b62d904ff4318bad8b65f9c5f131a"
CHOICE_REPARATION_OUI = "89WB4MJ"
Q_CONFIRMATION_REPRESENTANT = "22b7e585b6b64417aeb27af4588a2ea0"
CHOICE_CONFIRMATION_OUI = "psL681v"

LOG_FILE_NAME = "data_verification_maintenance_reguliere_log.xlsx"

LOG_HEADERS = [
    "Dimension", "Sous-dimension", "Response Code", "Water Point ID",
    "Description", "Champ concerné", "Détails", "Statut", "Première détection",
    "Dernière détection", "Date de résolution",
]

# Déploiement "Maintenance préventive régulière" du formulaire registre (lu dans la
# définition du formulaire le 13/10/2026).
DEPLOYMENT_LOG_VERIFICATION = "69d1692f3cbd424abdbb86c883b8fde4"
# Même énumérateur que pour Appel maintenance préventive et Réparation après panne.
ENUMERATEUR_LOG_VERIFICATION = "0336889caebc433d8ab0bba5dc919bed"

# Colonnes à résoudre dans le datagrid : clé interne -> libellé normalisé (voir normalize).
WANTED_COLUMNS = {
    "pump": "is the pump functionning properly",
    "start": "when did the problem start",
    "repair_possible": "is the repair possible during the preventive maintenance mission",
    "after_repair": "does the pump function properly after the repair",
    "completion": "completion date of the work",
    "signal": "signal code",
}

PUMP_VALUES = {"yes", "partly", "no"}
YES_NO_VALUES = {"yes", "no"}


# ---------------------------------------------------------------------------
# Colonnes du datagrid
# ---------------------------------------------------------------------------

def normalize(label):
    """Compare les libellés sans tenir compte de la casse, des espaces (dont insécables)
    ni du '?' final : le formulaire écrit 'properly?' là où le manuel écrit 'properly ?'."""
    text = (label or "").replace("\xa0", " ").lower()
    text = re.sub(r"\s+", " ", text).strip()
    return text.rstrip("?").strip()


def find_column(headers, wanted):
    for h in headers:
        if normalize(h) == wanted:
            return h
    for h in headers:
        if normalize(h).startswith(wanted):
            return h
    return None


def find_water_point_column(headers):
    for h in headers:
        n = normalize(h)
        if n.startswith("water point id") and n.endswith("unique id"):
            return h
    return None


def resolve_columns(rows):
    """Renvoie {clé: nom de colonne ou None}. Avertit pour chaque colonne introuvable."""
    headers = list(rows[0].keys()) if rows else []
    cols = {key: find_column(headers, wanted) for key, wanted in WANTED_COLUMNS.items()}
    cols["wp"] = find_water_point_column(headers)
    missing = [k for k, v in cols.items() if v is None]
    if missing:
        print(f"AVERTISSEMENT : colonnes introuvables dans le datagrid : {missing}. "
              f"Les contrôles qui en dépendent sont sautés. Colonnes disponibles : "
              f"{headers}", file=sys.stderr)
    return cols


def val(row, col):
    return (row.get(col) or "").strip() if col else ""


# ---------------------------------------------------------------------------
# Complétude : Water Point ID vide mais référence présente dans la réponse brute
# ---------------------------------------------------------------------------

def enrich_missing_water_point_ids(rows, client_id, wp_col):
    """Pour les réponses 'Final' dont le Water Point ID est vide dans le datagrid, retrouve
    le code réellement répondu dans la réponse brute (tableau `entities`) : site supprimé ou
    inaccessible plutôt que champ jamais rempli. Marque la ligne avec `_wp_enriched`."""
    if not wp_col:
        return 0
    candidates = [r for r in rows if r.get("Status") == "Final" and not val(r, wp_col)]
    if not candidates:
        return 0

    codes = sorted({r.get("Response Code", "") for r in candidates if r.get("Response Code")})
    raw_by_code = fetch_raw_responses_by_code(client_id, FORM_MAINTENANCE, codes)

    already_resolved_by_rc = defaultdict(set)
    for r in rows:
        rc, wp = r.get("Response Code", ""), val(r, wp_col)
        if rc and wp:
            already_resolved_by_rc[rc].add(wp)

    pools = {}
    for rc, raws in raw_by_code.items():
        already = already_resolved_by_rc.get(rc, set())
        pools[rc] = [raw for raw in raws if extract_water_point_code(raw) not in already]

    count = 0
    for r in candidates:
        pool = pools.get(r.get("Response Code", "")) or []
        if not pool:
            continue
        wp_code = extract_water_point_code(pool.pop(0))
        if wp_code:
            r[wp_col] = wp_code
            r["_wp_enriched"] = True
            already_resolved_by_rc[r.get("Response Code", "")].add(wp_code)
            count += 1
    return count


# ---------------------------------------------------------------------------
# Fiabilité : liste des points d'eau réhabilités avec succès
# ---------------------------------------------------------------------------

def build_success_list(client_id):
    """Ensemble des Water Point ID réhabilités avec succès : réponses 'Première
    réhabilitation' du formulaire combiné, au statut Final, avec « La réparation ou la
    maintenance a réussi ? » = Oui ET « confirmation du représentant local » = Oui.
    Passe par les réponses brutes (IDs de questions exacts) plutôt que par les colonnes du
    datagrid Réhabilitation, dont les libellés sont ambigus (deux versions de chaque
    question dans le formulaire)."""
    flt = {
        "form": FORM_COMBINE,
        "status": "final",
        f"data.{Q_TYPE_TRAVAUX}.value": CHOICE_PREMIERE_REHAB,
        f"data.{Q_REPARATION_REUSSIE}.value": CHOICE_REPARATION_OUI,
        f"data.{Q_CONFIRMATION_REPRESENTANT}.value": CHOICE_CONFIRMATION_OUI,
    }
    resp = requests.get(
        f"{MWATER_API_BASE}/responses",
        params={"client": client_id, "filter": json.dumps(flt), "limit": 10000},
        timeout=180,
    )
    raise_for_status_verbose(resp)

    success = set()
    for item in resp.json():
        d = item.get("data", {})
        # Re-vérification côté script : ne dépend pas de l'interprétation du filtre.
        if (item.get("status") != "final"
                or (d.get(Q_TYPE_TRAVAUX) or {}).get("value") != CHOICE_PREMIERE_REHAB
                or (d.get(Q_REPARATION_REUSSIE) or {}).get("value") != CHOICE_REPARATION_OUI
                or (d.get(Q_CONFIRMATION_REPRESENTANT) or {}).get("value")
                != CHOICE_CONFIRMATION_OUI):
            continue
        for entity in item.get("entities", []):
            if entity.get("entityType") == "water_point" and entity.get("value"):
                success.add(str(entity["value"]).strip())
    return success


# ---------------------------------------------------------------------------
# Anomalie
# ---------------------------------------------------------------------------

class Anomaly:
    def __init__(self, dimension, subdimension, response_code, water_point_id, description,
                 details="", champ_concerne=""):
        self.dimension = dimension
        self.subdimension = subdimension
        self.response_code = response_code
        self.water_point_id = water_point_id
        self.description = description
        self.details = details
        self.champ_concerne = champ_concerne


# ---------------------------------------------------------------------------
# Dimensions
# ---------------------------------------------------------------------------

def check_completude(rows, cols):
    anomalies = []
    wp_col = cols["wp"]
    for r in rows:
        status, rc, wp = r.get("Status", ""), r.get("Response Code", ""), val(r, wp_col)
        if status == "Draft":
            anomalies.append(Anomaly(
                "Complétude", "Brouillon non soumis", rc, wp, "Brouillon jamais soumis",
                f"Drafted On: {r.get('Drafted On', '')}"))
        elif status == "Final" and wp_col and not wp:
            anomalies.append(Anomaly(
                "Complétude", "Water Point ID manquant", rc, "",
                "Réponse finalisée sans Water Point ID renseigné"))
        elif status == "Final" and r.get("_wp_enriched"):
            anomalies.append(Anomaly(
                "Complétude", "Site associé supprimé/inaccessible", rc, wp,
                "Water Point ID répondu mais site associé supprimé ou inaccessible",
                f"Code référencé : {wp} — non résolu dans l'export mWater malgré une "
                f"réponse valide (à vérifier dans le portail)"))
    return anomalies


def check_promptitude(rows, cols):
    anomalies = []
    for r in rows:
        drafted, submitted = parse_dt(r.get("Drafted On")), parse_dt(r.get("Submitted On"))
        if drafted and submitted and submitted < drafted:
            anomalies.append(Anomaly(
                "Promptitude", "Chronologie Drafted/Submitted", r.get("Response Code", ""),
                val(r, cols["wp"]), "Submitted On antérieur à Drafted On",
                f"Drafted On: {drafted} / Submitted On: {submitted}"))
    return anomalies


def check_validite(rows, cols):
    anomalies = []
    for r in rows:
        rc, wp = r.get("Response Code", ""), val(r, cols["wp"])

        signal = val(r, cols["signal"])
        if signal and not SIGNAL_CODE_RE.match(signal):
            corrected, note = propose_correction(signal)
            description = "Signal code ne respecte pas le format attendu"
            if note:
                description += f" ({note})"
            anomalies.append(Anomaly(
                "Validité", "Format signal code", rc, wp, description,
                details=f"Signal code correct proposé : {corrected}" if corrected else "",
                champ_concerne=signal))

        if wp and not wp.isdigit():
            anomalies.append(Anomaly(
                "Validité", "Format Water Point ID", rc, wp, "Water Point ID non numérique"))

        for key, allowed, label in (
            ("pump", PUMP_VALUES, "Is the pump functionning properly ?"),
            ("repair_possible", YES_NO_VALUES,
             "Is the repair possible during the preventive maintenance mission ?"),
            ("after_repair", YES_NO_VALUES, "Does the pump function properly after the repair ?"),
        ):
            value = val(r, cols[key])
            if value and value.lower() not in allowed:
                anomalies.append(Anomaly(
                    "Validité", f"Valeur '{label}'", rc, wp,
                    "Hors des valeurs attendues", champ_concerne=value))
    return anomalies


def check_unicite(rows, cols):
    """Clé = signal code parmi les réponses non-brouillon ; un signal code vide n'est
    jamais un doublon ; un rejet 'doublon' ne compte que si la réponse est toujours
    Rejected."""
    anomalies = []
    counts = Counter(
        val(r, cols["signal"]) for r in rows
        if r.get("Status") != "Draft" and val(r, cols["signal"])
    )
    duplicated = {code for code, n in counts.items() if n > 1}
    for r in rows:
        signal = val(r, cols["signal"])
        rc, wp = r.get("Response Code", ""), val(r, cols["wp"])
        if r.get("Status") != "Draft" and signal in duplicated:
            anomalies.append(Anomaly(
                "Unicité", "Doublon signal code", rc, wp, "Signal code dupliqué",
                champ_concerne=signal))
        elif r.get("Status") == "Rejected" and "doublon" in (r.get("Rejection message") or "").lower():
            anomalies.append(Anomaly(
                "Unicité", "Doublon signal code", rc, wp,
                "Rejet mWater pour doublon de signal code", champ_concerne=signal))
    return anomalies


def check_coherence(rows, cols):
    """Cohérence interne au formulaire (champs conditionnels). Les règles croisées avec
    Appel ne sont pas refaites ici."""
    anomalies = []
    for r in rows:
        if r.get("Status") == "Draft":
            continue
        rc, wp = r.get("Response Code", ""), val(r, cols["wp"])
        pump = val(r, cols["pump"]).lower()
        possible = val(r, cols["repair_possible"]).lower()
        after = val(r, cols["after_repair"])
        start_raw = val(r, cols["start"])
        sub = "Champs conditionnels incohérents"

        if cols["pump"] and pump in ("partly", "no"):
            if cols["start"] and not start_raw:
                anomalies.append(Anomaly(
                    "Cohérence", sub, rc, wp,
                    f"Pompe '{pump}' mais 'When did the problem start ?' n'est pas renseigné"))
            if cols["repair_possible"] and not possible:
                anomalies.append(Anomaly(
                    "Cohérence", sub, rc, wp,
                    f"Pompe '{pump}' mais 'Is the repair possible during the preventive "
                    f"maintenance mission ?' n'est pas renseigné"))

        if cols["repair_possible"] and cols["after_repair"] and possible == "yes" and not after:
            anomalies.append(Anomaly(
                "Cohérence", sub, rc, wp,
                "Réparation possible = Yes mais 'Does the pump function properly after the "
                "repair ?' n'est pas renseigné"))

        start, completion = parse_dt(start_raw), parse_dt(val(r, cols["completion"]))
        if start and completion and start.date() > completion.date():
            anomalies.append(Anomaly(
                "Cohérence", "Dates incohérentes", rc, wp,
                "'When did the problem start ?' postérieur à 'Completion date of the work'",
                f"Début du problème : {start.date()} / Fin des travaux : {completion.date()}"))
    return anomalies


def check_fiabilite(rows, cols, success_wp_ids):
    """Le Water Point ID doit appartenir à la liste des points d'eau réhabilités avec
    succès. Un ID absent de la liste est une anomalie (réponses non brouillon, non
    rejetées) ; un rejet « ID incorrect » ne compte que s'il est toujours actif
    (Status = Rejected)."""
    anomalies = []
    wp_col = cols["wp"]
    if not wp_col:
        return anomalies
    for r in rows:
        status, rc, wp = r.get("Status", ""), r.get("Response Code", ""), val(r, wp_col)
        if not wp or status == "Draft":
            continue
        rejection = (r.get("Rejection message") or "").lower()
        if status == "Rejected":
            if "id incorrect" in rejection:
                dans_liste = "présent" if wp in success_wp_ids else "absent"
                anomalies.append(Anomaly(
                    "Fiabilité", "Rejet ID incorrect non résolu", rc, wp,
                    "Rejet 'ID incorrect' toujours actif",
                    f"Water Point ID {dans_liste} de la liste des réhabilitations réussies"))
            continue
        if wp not in success_wp_ids:
            anomalies.append(Anomaly(
                "Fiabilité", "Point d'eau absent des réhabilitations réussies", rc, wp,
                "Water Point ID absent de la liste des points d'eau réhabilités avec succès"))
    return anomalies


# ---------------------------------------------------------------------------
# Log Excel — Nouveau / Toujours ouvert / Résolu / Supprimé
# ---------------------------------------------------------------------------

def anomaly_key(a):
    return (a.dimension, a.subdimension, a.response_code, a.description, a.water_point_id)


def build_submitted_on_lookup(rows):
    lookup = {}
    for r in rows:
        dt = parse_dt(r.get("Submitted On"))
        if r.get("Response Code") and dt:
            lookup[r["Response Code"]] = dt.strftime("%d/%m/%Y")
    return lookup


def merge_with_log(current_anomalies, existing_rows, today_str, submitted_on_lookup=None,
                   existing_response_codes=None):
    """Même logique à 4 statuts que verify_maintenance_preventive.py (voir sa docstring).
    `existing_rows` vient de download_existing_log, qui normalise les cellules vides en ""
    (sans quoi les anomalies sans point d'eau ne seraient jamais retrouvées)."""
    submitted_on_lookup = submitted_on_lookup or {}
    existing_response_codes = existing_response_codes or set()
    existing_open, already_closed = {}, []
    for row in existing_rows:
        key = (row.get("Dimension"), row.get("Sous-dimension"), row.get("Response Code"),
               row.get("Description"), row.get("Water Point ID"))
        if row.get("Statut") in ("Résolu", "Supprimé"):
            already_closed.append(row)
        else:
            existing_open[key] = row

    merged, seen_keys, new_count = [], set(), 0
    for a in current_anomalies:
        key = anomaly_key(a)
        if key in seen_keys:
            continue  # même anomalie détectée deux fois dans cette exécution
        seen_keys.add(key)
        existing = existing_open.get(key)
        if not existing:
            new_count += 1
        merged.append({
            "Dimension": a.dimension,
            "Sous-dimension": a.subdimension,
            "Response Code": a.response_code,
            "Water Point ID": a.water_point_id,
            "Description": a.description,
            "Champ concerné": a.champ_concerne,
            "Détails": a.details,
            "Statut": "Toujours ouvert" if existing else "Nouveau",
            "Première détection": existing.get("Première détection") if existing else today_str,
            "Dernière détection": today_str,
            "Date de résolution": "",
        })

    resolved_count = deleted_count = 0
    for key, row in existing_open.items():
        if key in seen_keys:
            continue
        row = dict(row)
        if row.get("Response Code") not in existing_response_codes:
            row["Statut"] = "Supprimé"
            row["Date de résolution"] = row.get("Première détection", "")
            deleted_count += 1
        else:
            row["Statut"] = "Résolu"
            row["Date de résolution"] = submitted_on_lookup.get(row.get("Response Code"), today_str)
            resolved_count += 1
        merged.append(row)

    merged.extend(already_closed)
    return merged, new_count, resolved_count, deleted_count


def build_log_report(merged_rows, output_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Anomalies"
    ws.append(LOG_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in merged_rows:
        ws.append([row.get(h, "") for h in LOG_HEADERS])

    open_rows = [r for r in merged_rows if r.get("Statut") not in ("Résolu", "Supprimé")]
    summary = wb.create_sheet("Résumé")
    summary.append(["Dimension", "Nombre d'anomalies ouvertes"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    counts = Counter(r.get("Dimension") for r in open_rows)
    for dimension in ["Complétude", "Promptitude", "Validité", "Unicité", "Cohérence", "Fiabilité"]:
        summary.append([dimension, counts.get(dimension, 0)])
    summary.append(["Total ouvert", len(open_rows)])
    wb.save(output_path)


def send_confirmation_email(token, sender, recipients, file_name, open_count, new_count,
                            resolved_count, deleted_count, counts_by_dimension):
    lines = "".join(f"<li>{dim} : {n}</li>" for dim, n in counts_by_dimension.items())
    body_html = (
        f"<p>Vérification de données \"Maintenance préventive régulière\" exécutée.</p>"
        f"<p>Anomalies actuellement ouvertes : <b>{open_count}</b> "
        f"(dont {new_count} nouvelles) — {resolved_count} résolue(s) et {deleted_count} "
        f"supprimée(s) (réponse mWater disparue) depuis la dernière exécution.</p>"
        f"<ul>{lines}</ul>"
        f"<p>Log mis à jour sur SharePoint : <b>{file_name}</b></p>"
    )
    send_html_email(token, sender, recipients, f"Vérification de données — {file_name}", body_html)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    mwater_username = os.environ["MWATER_USERNAME"].strip()
    mwater_password = os.environ["MWATER_PASSWORD"].strip()
    azure_tenant_id = os.environ["AZURE_TENANT_ID"].strip()
    azure_client_id = os.environ["AZURE_CLIENT_ID"].strip()
    azure_client_secret = os.environ["AZURE_CLIENT_SECRET"].strip()
    sharepoint_folder_link = os.environ["SHAREPOINT_FOLDER_LINK"].strip()
    email_sender = os.environ["EMAIL_SENDER"].strip()
    email_recipients = os.environ["EMAIL_RECIPIENTS"].strip()

    print("Authentification mWater...")
    client_id = mwater_login(mwater_username, mwater_password)

    print("Téléchargement du datagrid Maintenance préventive...")
    rows = download_datagrid(DATAGRID_MAINTENANCE, client_id)
    print(f"  {len(rows)} réponses")
    cols = resolve_columns(rows)

    print("Résolution des Water Point ID vides (site supprimé mais réponse valide)...")
    enriched = enrich_missing_water_point_ids(rows, client_id, cols["wp"])
    print(f"  {enriched} lignes enrichies avec leur code brut")

    print("Construction de la liste des points d'eau réhabilités avec succès...")
    success_wp_ids = build_success_list(client_id)
    print(f"  {len(success_wp_ids)} points d'eau (≈725 attendus)")

    print("Application des règles de vérification (6 dimensions)...")
    anomalies = []
    anomalies += check_completude(rows, cols)
    anomalies += check_promptitude(rows, cols)
    anomalies += check_validite(rows, cols)
    anomalies += check_unicite(rows, cols)
    anomalies += check_coherence(rows, cols)
    if success_wp_ids:
        anomalies += check_fiabilite(rows, cols, success_wp_ids)
    else:
        print("AVERTISSEMENT : liste des réhabilitations réussies vide, dimension "
              "Fiabilité sautée.", file=sys.stderr)
    print(f"  {len(anomalies)} anomalies détectées")

    print("Authentification Microsoft Graph...")
    token = graph_token(azure_tenant_id, azure_client_id, azure_client_secret)
    drive_id, folder_item_id = resolve_share_link(token, sharepoint_folder_link)

    print("Téléchargement du log existant sur SharePoint...")
    existing_rows = download_existing_log(token, drive_id, folder_item_id, LOG_FILE_NAME)
    print(f"  {len(existing_rows)} lignes déjà présentes dans le log")

    today_str = datetime.now().strftime("%d/%m/%Y")
    merged_rows, new_count, resolved_count, deleted_count = merge_with_log(
        anomalies, existing_rows, today_str,
        submitted_on_lookup=build_submitted_on_lookup(rows),
        existing_response_codes={r.get("Response Code", "") for r in rows})
    open_rows = [r for r in merged_rows if r.get("Statut") not in ("Résolu", "Supprimé")]
    print(f"  {len(open_rows)} anomalies ouvertes ({new_count} nouvelles, {resolved_count} "
          f"résolues, {deleted_count} supprimées)")

    try:
        insert_log_in_mwater(merged_rows, client_id, DEPLOYMENT_LOG_VERIFICATION,
                             ENUMERATEUR_LOG_VERIFICATION)
    except Exception as exc:
        print(f"AVERTISSEMENT : échec de l'écriture dans le registre mWater ({exc}) — "
              f"le fichier Excel reste la source fiable pour cette exécution.", file=sys.stderr)

    output_path = f"/tmp/{LOG_FILE_NAME}"
    build_log_report(merged_rows, output_path)
    print("Dépôt du log mis à jour sur SharePoint...")
    upload_to_sharepoint(token, drive_id, folder_item_id, output_path, LOG_FILE_NAME)

    print("Envoi de l'email de confirmation...")
    send_confirmation_email(token, email_sender, email_recipients, LOG_FILE_NAME,
                            len(open_rows), new_count, resolved_count, deleted_count,
                            Counter(r.get("Dimension") for r in open_rows))
    print("Terminé.")


if __name__ == "__main__":
    main()
