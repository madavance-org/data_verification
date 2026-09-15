"""Microsoft Graph (SharePoint + email), partagé par tous les scripts de ce repo."""

import base64
import io

import requests

from common.http_utils import raise_for_status_verbose

GRAPH_API_BASE = "https://graph.microsoft.com/v1.0"


def graph_token(tenant_id, client_id, client_secret):
    resp = requests.post(
        f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token",
        data={
            "grant_type": "client_credentials",
            "client_id": client_id,
            "client_secret": client_secret,
            "scope": "https://graph.microsoft.com/.default",
        },
        timeout=30,
    )
    raise_for_status_verbose(resp)
    return resp.json()["access_token"]


def resolve_share_link(token, share_url):
    """Résout un lien de partage SharePoint (ex. https://.../:f:/s/...) en (drive_id, item_id)
    via l'API Graph /shares/{shareId}/driveItem. Évite d'avoir à extraire les IDs à la main :
    il suffit de coller le lien de partage tel quel dans le secret SHAREPOINT_FOLDER_LINK."""
    b64 = base64.b64encode(share_url.encode("utf-8")).decode("utf-8")
    encoded = "u!" + b64.replace("+", "-").replace("/", "_").rstrip("=")
    resp = requests.get(
        f"{GRAPH_API_BASE}/shares/{encoded}/driveItem",
        headers={"Authorization": f"Bearer {token}"},
        timeout=30,
    )
    raise_for_status_verbose(resp)
    data = resp.json()
    return data["parentReference"]["driveId"], data["id"]


def find_child_item_id(token, drive_id, folder_item_id, file_name):
    """Cherche un fichier par nom parmi les enfants d'un dossier, en adressage 100% par ID
    (pas de syntaxe 'items/{id}:/{path}:' — celle-ci renvoie un 400 'Invalid request' pour
    certains dossiers en authentification applicative, alors qu'elle fonctionne en délégué ;
    cause probable : résolution par chemin qui échoue côté Graph pour cet item, indépendamment
    des permissions elles-mêmes). Retourne l'ID de l'item trouvé, ou None."""
    resp = requests.get(
        f"{GRAPH_API_BASE}/drives/{drive_id}/items/{folder_item_id}/children",
        headers={"Authorization": f"Bearer {token}"},
        timeout=60,
    )
    raise_for_status_verbose(resp)
    for item in resp.json().get("value", []):
        if item.get("name") == file_name:
            return item.get("id")
    return None


def download_existing_log(token, drive_id, folder_item_id, file_name,
                           sheet_name="Anomalies", key_column="Dimension", row_hook=None):
    """Télécharge un log Excel existant sur SharePoint et le retourne comme liste de dict (une
    par ligne non vide). Retourne [] si le fichier n'existe pas encore (première exécution).

    `key_column` sert à ignorer les lignes vides éventuelles — varie selon le schéma du log
    ('Dimension' pour un script à 6 dimensions, 'Formulaire' pour un contrôle ponctuel par
    formulaire). `row_hook`, si fourni, est appliqué à chaque ligne lue (ex. migration d'un
    ancien format de log — voir migrer_ligne_champ_concerne dans verify_maintenance_preventive.py,
    qui reste spécifique à ce script et n'est pas partagée ici).

    Normalise les cellules vides (None, tel que renvoyé par openpyxl) en chaîne vide "" :
    sans ça, une clé de suivi construite à partir d'une anomalie fraîchement détectée (qui
    vaut toujours "" pour un champ vide, ex. Water Point ID manquant) ne correspond jamais à
    la même ligne relue depuis ce fichier (qui vaudrait None) — l'anomalie n'est alors jamais
    reconnue comme "toujours ouverte" d'une exécution à l'autre, et le log recrée une nouvelle
    ligne "Nouveau" à chaque fois. Bug réel constaté le 15/09/2026 sur des anomalies sans point
    d'eau (jusqu'à 38 lignes dupliquées pour une seule anomalie sur ~2 mois)."""
    item_id = find_child_item_id(token, drive_id, folder_item_id, file_name)
    if not item_id:
        return []

    resp = requests.get(
        f"{GRAPH_API_BASE}/drives/{drive_id}/items/{item_id}/content",
        headers={"Authorization": f"Bearer {token}"},
        timeout=60,
    )
    raise_for_status_verbose(resp)

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb[sheet_name]
    headers = [c.value for c in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {h: ("" if v is None else v) for h, v in zip(headers, values)}
        if row.get(key_column):
            rows.append(row_hook(row) if row_hook else row)
    return rows


def upload_to_sharepoint(token, drive_id, folder_item_id, file_path, file_name,
                          content_type="application/octet-stream"):
    """Dépose/écrase le fichier, en adressage 100% par ID (voir note dans find_child_item_id).
    S'il n'existe pas encore, on crée d'abord un fichier vide dans le dossier
    (POST .../children), puis on écrit son contenu par ID."""
    with open(file_path, "rb") as f:
        content = f.read()

    item_id = find_child_item_id(token, drive_id, folder_item_id, file_name)
    if not item_id:
        resp = requests.post(
            f"{GRAPH_API_BASE}/drives/{drive_id}/items/{folder_item_id}/children",
            headers={"Authorization": f"Bearer {token}"},
            json={
                "name": file_name,
                "file": {},
                "@microsoft.graph.conflictBehavior": "replace",
            },
            timeout=60,
        )
        raise_for_status_verbose(resp)
        item_id = resp.json()["id"]

    resp = requests.put(
        f"{GRAPH_API_BASE}/drives/{drive_id}/items/{item_id}/content",
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": content_type,
        },
        data=content,
        timeout=120,
    )
    raise_for_status_verbose(resp)
    return resp.json()


def send_html_email(token, sender, recipients, subject, body_html):
    """Envoie un email HTML via Microsoft Graph (sendMail), au nom de `sender`. `recipients`
    est une chaîne d'adresses séparées par des virgules."""
    resp = requests.post(
        f"{GRAPH_API_BASE}/users/{sender}/sendMail",
        headers={"Authorization": f"Bearer {token}"},
        json={
            "message": {
                "subject": subject,
                "body": {"contentType": "HTML", "content": body_html},
                "toRecipients": [{"emailAddress": {"address": addr.strip()}} for addr in recipients.split(",")],
            }
        },
        timeout=30,
    )
    raise_for_status_verbose(resp)
